"""Advance Voucher, Advance Adjustment Voucher and Cash Voucher – filled into the original Excel
files with openpyxl (works on Windows and on the Linux server), then exported to PDF by convert.py
(Excel on Windows, LibreOffice on the server)."""
import os
from copy import copy
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.cell.cell import MergedCell
from openpyxl.styles.cell_style import StyleArray

from convert import xlsx_to_pdf
from gen_word import parse_date, fy
from money import words, r2

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "doc_templates")

# Indian grouping that both Excel and LibreOffice understand
INR0 = '[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0'
INR2 = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
DATE = "dd-mm-yyyy"


# ------------------------------------------------------------------ sheet helpers
def _date(ws, ref, d, fmt=DATE):
    d = parse_date(d)
    ws[ref] = datetime(d.year, d.month, d.day)
    ws[ref].number_format = fmt


def _shift_rows(ws, from_row, delta):
    """Move every row >= from_row by delta (negative = up), with merges and row heights.
    Rows that would be overwritten when moving up (from_row+delta .. from_row-1) are discarded."""
    if delta == 0:
        return
    max_row, max_col = ws.max_row, ws.max_column
    if delta < 0:
        gone_lo, gone_hi = from_row + delta, from_row - 1
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row <= gone_hi and rng.max_row >= gone_lo:
                ws.unmerge_cells(str(rng))
        for r in range(gone_lo, gone_hi + 1):
            for c in range(1, max_col + 1):
                ws.cell(r, c).value = None
    moved = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= from_row:
            moved.append(CellRange(min_col=rng.min_col, min_row=rng.min_row + delta,
                                   max_col=rng.max_col, max_row=rng.max_row + delta))
            ws.unmerge_cells(str(rng))
    if from_row <= max_row:
        ws.move_range(f"A{from_row}:{ws.cell(1, max_col).column_letter}{max_row}", rows=delta)
    heights = {r: ws.row_dimensions[r].height for r in range(from_row, max_row + 1)}
    for r in range(from_row, max_row + 1):
        ws.row_dimensions[r].height = None
    for r, hgt in heights.items():
        ws.row_dimensions[r + delta].height = hgt
    if delta < 0:  # wipe the rows left behind at the bottom
        for r in range(max_row + delta + 1, max_row + 1):
            ws.row_dimensions[r].height = None
            for c in range(1, max_col + 1):
                cell = ws.cell(r, c)
                if not isinstance(cell, MergedCell):
                    cell.value = None
                    cell._style = StyleArray()
    for rng in moved:
        ws.merge_cells(rng.coord)


def _copy_row_style(ws, src, dst, cols):
    for c in cols:
        s, d = ws.cell(src, c), ws.cell(dst, c)
        d._style = copy(s._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def _fit_rows(ws, first, n_template, n_needed, cols, merge=None):
    """Grow/shrink an item block of n_template rows starting at `first` to n_needed rows.
    Every row gets the first row's formatting; `merge` = (col_from, col_to) merged per row."""
    n = max(1, n_needed)
    if n > n_template:
        _shift_rows(ws, first + n_template, n - n_template)
    elif n < n_template:
        _shift_rows(ws, first + n_template, n - n_template)
    for r in range(first, first + n):
        for c in cols:
            cell = ws.cell(r, c)
            if not isinstance(cell, MergedCell):
                cell.value = None
        if r != first:
            _copy_row_style(ws, first, r, cols)
        if merge:
            ref = f"{merge[0]}{r}:{merge[1]}{r}"
            if not any(str(m) == ref for m in ws.merged_cells.ranges):
                ws.merge_cells(ref)
    return n


def _one_page(ws):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def _ticks(ws):
    """The original uses Excel's new in-cell checkboxes, which other programs print as '###'."""
    for a in ("A1", "A2"):
        ws[a] = "☑"
        f = copy(ws[a].font)
        f.name = "Segoe UI Symbol"
        ws[a].font = f


def _set(ws, ref, value):
    """Write into a cell; if it is inside a merged block, write into its top-left cell."""
    for rng in ws.merged_cells.ranges:
        if ref in rng:
            ref = rng.start_cell.coordinate
            break
    ws[ref] = value
    return ws[ref]


# ------------------------------------------------------------------ advance request
def build_advance_voucher(adv, profile, budget, xlsx_out, pdf_out):
    """adv: {date, amount, items:[{description, qty, rate}], purpose}"""
    wb = load_workbook(os.path.join(TPL, "advance_voucher.xlsx"))
    ws = wb["Sheet1"]
    items = adv.get("items") or []
    amount = r2(adv.get("amount") or sum(float(i.get("qty") or 1) * float(i.get("rate") or 0) for i in items))
    d = parse_date(adv.get("date"))

    _ticks(ws)
    _date(ws, "H6", d)
    _set(ws, "C8", profile.get("payee", ""))
    _set(ws, "H8", profile.get("contact", ""))
    _set(ws, "B9", profile.get("payee_dept", ""))
    _set(ws, "G9", amount)
    _set(ws, "B10", words(amount))
    _set(ws, "C11", adv.get("purpose") or profile.get("advance_purpose", ""))
    from datetime import timedelta
    _date(ws, "G13", d + timedelta(days=30))
    _set(ws, "D22", profile.get("project", ""))
    _set(ws, "D23", profile.get("budget_code", ""))
    prov, used = r2(budget.get("provision", 0)), r2(budget.get("utilized", 0))
    for ref, v in (("E24", prov), ("E25", used), ("E26", amount), ("E27", r2(prov - used - amount))):
        _set(ws, ref, v).number_format = INR0
    _set(ws, "E29", profile.get("advance_approver", "HOD,EC"))

    n = _fit_rows(ws, 34, 10, len(items), list(range(2, 9)), merge=("B", "E"))
    total = 0.0
    for i, it in enumerate(items):
        r = 34 + i
        q, rate = it.get("qty"), it.get("rate")
        ws[f"B{r}"] = it.get("description", "")
        line = float(q or 1) * float(rate or 0)
        total += line
        if q not in (None, ""):
            ws[f"F{r}"] = q
            ws[f"G{r}"] = rate
        ws[f"H{r}"] = line
    ws[f"H{34 + n}"] = total
    _one_page(ws)
    wb.save(xlsx_out)
    return xlsx_to_pdf(xlsx_out, pdf_out)


# ------------------------------------------------------------------ advance adjustment
def build_advance_adjustment(adj, profile, budget, xlsx_out, pdf_out):
    """adj: {date, bills:[{biller, bill_no, bill_date, gst, items, amount}], remaining_adjustment}"""
    wb = load_workbook(os.path.join(TPL, "advance_adjustment.xlsx"))
    ws = wb["Sheet2"]
    for other in [s for s in wb.worksheets if s.title != "Sheet2"]:
        wb.remove(other)
    bills = adj.get("bills") or []
    total = r2(sum(float(b.get("amount") or 0) for b in bills))
    rounded = round(total)

    _date(ws, "B2", adj.get("date"))
    _set(ws, "C4", profile.get("payee", ""))
    _set(ws, "F4", f"Contact: {profile.get('contact', '')}")
    _set(ws, "A7", adj.get("particular") or profile.get("adjust_particular", ""))
    _set(ws, "F7", total)
    _set(ws, "B13", words(rounded))
    _set(ws, "F13", rounded)
    _set(ws, "B18", profile.get("voucher_mentor", ""))
    _set(ws, "C18", profile.get("voucher_hod", ""))
    _set(ws, "D21", profile.get("project", ""))
    _set(ws, "D22", profile.get("budget_code", ""))
    prov, used = r2(budget.get("provision", 0)), r2(budget.get("utilized", 0))
    remaining = r2(adj.get("remaining_adjustment") or 0)
    _set(ws, "E23", prov)
    _set(ws, "E24", used)
    _set(ws, "E25", rounded)
    _set(ws, "E26", remaining)
    _set(ws, "E27", r2(prov - (used + remaining + rounded)))
    _set(ws, "C29", f"Name of the Approving Authority : {profile.get('voucher_approver', '')}")
    code = fy(adj.get("date")).replace("-", "")[2:]  # 2026-27 -> 2627
    prefix = (profile.get("adjust_ref_prefix") or "NU/IT/{project}/{fy}/Advance Adjustment Voucher/") \
        .replace("{project}", profile.get("project", "")).replace("{fy}", code)
    _set(ws, "A31", prefix + profile.get("payee", ""))

    # bills table starts at row 34; clear what the template had below, then write ours
    for r in range(34, max(ws.max_row, 34) + 1):
        for c in range(1, 8):
            if not isinstance(ws.cell(r, c), MergedCell):
                ws.cell(r, c).value = None
    n = max(1, len(bills))
    for r in range(35, 35 + max(n, 40)):
        if r < 34 + n:
            _copy_row_style(ws, 34, r, range(1, 8))
        else:
            for c in range(1, 8):
                ws.cell(r, c)._style = StyleArray()
    for i, b in enumerate(bills):
        r = 34 + i
        ws[f"A{r}"] = i + 1
        ws[f"B{r}"] = b.get("biller", "")
        ws[f"C{r}"] = str(b.get("bill_no", ""))
        if b.get("bill_date"):
            _date(ws, f"D{r}", b["bill_date"])
        ws[f"E{r}"] = r2(b.get("gst") or 0)
        ws[f"F{r}"] = b.get("items", "")
        ws[f"G{r}"] = r2(b.get("amount") or 0)
    _set(ws, "G32", total)
    ws.print_area = f"A1:G{max(40, 33 + n)}"
    _one_page(ws)
    wb.save(xlsx_out)
    return xlsx_to_pdf(xlsx_out, pdf_out)


# ------------------------------------------------------------------ cash voucher (reimbursement)
def build_cash_voucher(cv, profile, budget, xlsx_out, pdf_out):
    """Reimbursement of a personal-fund purchase. No official template was provided, so this uses the
    Nirma 'Voucher for Advance' sheet (same header, logo and budget block) re-titled as a Cash Voucher.
    cv: {date, payee?, contact?, purpose, permission_ref, bills:[{biller, bill_no, bill_date, items, amount}]}"""
    wb = load_workbook(os.path.join(TPL, "advance_voucher.xlsx"))
    ws = wb["Sheet1"]
    bills = cv.get("bills") or []
    total = r2(sum(float(b.get("amount") or 0) for b in bills))

    _ticks(ws)
    _set(ws, "D5", "CASH VOUCHER")
    _date(ws, "H6", cv.get("date"))
    _set(ws, "C8", cv.get("payee") or profile.get("payee", ""))
    _set(ws, "H8", cv.get("contact") or profile.get("contact", ""))
    _set(ws, "B9", profile.get("payee_dept", ""))
    _set(ws, "F9", "Amount Rs.")
    _set(ws, "G9", total)
    _set(ws, "B10", words(round(total)))
    _set(ws, "A11", "Being payment of")
    c11 = _set(ws, "C11", cv.get("purpose") or (
        f"Reimbursement of material purchased for {profile.get('project', '')} ({profile.get('grant', '')})"))
    al = copy(c11.alignment)
    al.shrink_to_fit = True
    c11.alignment = al
    _set(ws, "A13", "Purchase Permission Ref. / Date")
    _set(ws, "G13", cv.get("permission_ref", ""))
    _set(ws, "A14", "Paid from personal funds; original bill(s) attached.")
    for col in "ABCDEFGHI":
        ws[f"{col}15"].value = None
    _set(ws, "D22", profile.get("project", ""))
    _set(ws, "D23", profile.get("budget_code", ""))
    prov, used = r2(budget.get("provision", 0)), r2(budget.get("utilized", 0))
    for ref, v in (("E24", prov), ("E25", used), ("E26", total), ("E27", r2(prov - used - total))):
        _set(ws, ref, v).number_format = INR2
    _set(ws, "E29", cv.get("approver") or profile.get("advance_approver", "HOD,EC"))
    _set(ws, "A32", "Details of bills attached")
    _set(ws, "B33", "Biller / Items")
    _set(ws, "F33", "Bill No.")
    _set(ws, "G33", "Bill Date")
    _set(ws, "H33", "Amount")

    n = _fit_rows(ws, 34, 10, len(bills), list(range(2, 9)), merge=("B", "E"))
    for i, b in enumerate(bills):
        r = 34 + i
        desc = b.get("biller", "")
        if b.get("items"):
            desc = f"{desc} – {b['items']}" if desc else b["items"]
        ws[f"B{r}"] = desc
        ws[f"F{r}"] = str(b.get("bill_no", ""))
        if b.get("bill_date"):
            _date(ws, f"G{r}", b["bill_date"], "dd-mm-yy")
            al = copy(ws[f"G{r}"].alignment)
            al.shrink_to_fit = True
            ws[f"G{r}"].alignment = al
        ws[f"H{r}"] = r2(b.get("amount") or 0)
    ws[f"H{34 + n}"] = total
    _one_page(ws)
    wb.save(xlsx_out)
    return xlsx_to_pdf(xlsx_out, pdf_out)
