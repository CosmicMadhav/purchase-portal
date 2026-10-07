"""Builds every document type and checks what is printed in it."""
import os
import re

import pytest
from docx import Document
from openpyxl import load_workbook

from conftest import quote
from docbuild import unique_cells
from gen_word import build_permission, build_po, build_rfq, quote_totals
from gen_excel import build_advance_voucher, build_advance_adjustment, build_cash_voucher
import gen_excel
import store

BUDGET = {"provision": 300000, "utilized": 259572.10, "available": 40427.90, "as_of": "2026-10-07"}


@pytest.fixture(autouse=True)
def no_pdf(monkeypatch):
    """Excel → PDF is tested separately (slow); here only the filled .xlsx is checked."""
    monkeypatch.setattr(gen_excel, "xlsx_to_pdf", lambda x, p: p)


def text(path):
    d = Document(path)
    return "\n".join(p.text for p in d.paragraphs)


def table_rows(path, idx):
    t = Document(path).tables[idx]
    return [[c.text.strip() for c in unique_cells(r)] for r in t.rows]


def money(s):
    m = re.search(r"-?[\d,]+\.\d\d", s)
    return float(m.group(0).replace(",", "")) if m else None


def three_quotes(rates=(1000, 1100, 1200), qty=4):
    return [quote(v, [("Motor driver", qty, r)]) for v, r in zip(["ALPHA", "BETA", "GAMMA"], rates)]


# ---------------------------------------------------------------- permission variants
def test_small_purchase_goes_to_hod_without_statement(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "Mechanical Component (Wheel hub)",
         "quotes": [quote("HYDRO SPRAY TECH", [("Wheel hub", 4, 625)])]}
    out = build_permission(p, profile, str(tmp_path / "a.docx"))
    t = text(out)
    assert "Subject: Permission to purchase for Mechanical Component (Wheel hub)" in t
    assert "To" in t and "HoD- EC (N.P Gajjar)" in t
    assert "Through" not in t and "Statement" not in t
    rows = table_rows(out, 0)
    assert rows[0][2].startswith("Unit Rate")                     # incl. tax format
    assert rows[1][2] == "Rs. 737.50" and rows[1][4] == "Rs. 2,950.00"
    assert rows[-1] == ["Grand Total(approx)", "Rs. 2,950.00"]
    assert len(Document(out).tables) == 1


def test_3k_to_10k_with_comparison_stays_with_hod(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "Motor drivers", "quotes": three_quotes((1500, 1700, 1600)), "selected": 0}
    t = text(build_permission(p, profile, str(tmp_path / "b.docx")))
    assert "Comparative Statement" in t and "It is recommended to purchase the item from ALPHA." in t
    assert "Through" not in t and "Director" not in t


def test_director_route_and_supplier_order(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "Motor drivers", "quotes": three_quotes((5000, 4000, 6000)), "selected": 1}
    out = build_permission(p, profile, str(tmp_path / "c.docx"))
    t = text(out)
    assert "Through" in t and "Asst. Registrar, SoT-NU" in t and "Director, SoT NU" in t
    assert "Vice President" not in t
    comp = table_rows(out, 1)
    assert comp[0] == ["Supplier name", "BETA", "ALPHA", "GAMMA"]   # selected vendor first
    assert "from BETA." in t


def test_vp_route_above_50k(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "Battery", "quotes": three_quotes((15000, 16000, 17000))}
    t = text(build_permission(p, profile, str(tmp_path / "d.docx")))
    assert "Exec. Registrar, NU" in t and "Vice President, NU" in t


def test_authority_override(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "X", "quotes": three_quotes((5000, 6000, 7000)), "authority_override": "vp"}
    assert "Vice President, NU" in text(build_permission(p, profile, str(tmp_path / "e.docx")))


def test_single_quote_statement_wording(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "IMU", "quotes": [quote("ALL TECH", [("IMU", 1, 2900)])], "authority_override": "director"}
    out = build_permission(p, profile, str(tmp_path / "f.docx"))
    t = text(out)
    assert " Statement" in t and "Comparative" not in t
    assert "It is requested to purchase the item from ALL TECH." in t
    assert len(table_rows(out, 1)[0]) == 2


def test_totals_printed_match_and_add_up(tmp_path, profile):
    q = quote("ALPHA", [("A", 3, 1234.56), ("B", 7, 99.99)], discount=100, other=180, other_gst=18, round_off=-0.3)
    p = {"date": "2026-10-07", "title": "Parts", "quotes": [q, quote("B", [("A", 1, 9000)]), quote("C", [("A", 1, 9500)])]}
    out = build_permission(p, profile, str(tmp_path / "g.docx"))
    t = quote_totals(q)
    rows = {r[0]: money(r[1]) for r in table_rows(out, 0) if len(r) == 2}
    assert rows["Total (₹)"] == t["subtotal"] and rows["Discount"] == t["discount"]
    assert rows["GST"] == t["gst"] and rows["Other charges (Round off)"] == t["other"]
    assert rows["Grand Total"] == t["grand"] == round(t["subtotal"] - t["discount"] + t["gst"] + t["other"], 2)
    assert f"Total cost of purchase: ₹ {t['grand']:,.2f}".replace(",", "") in text(out).replace(",", "")
    comp = table_rows(out, 1)
    grand_row = next(r for r in comp if r[0] == "Grand Total")
    assert money(grand_row[1]) == t["grand"]


def test_multi_item_subject_and_unequal_quotes(tmp_path, profile):
    q1 = quote("ALPHA", [("Jetson", 1, 24950), ("Case", 1, 1400), ("SSD", 1, 4500)])
    q2 = quote("BETA", [("Jetson kit", 1, 27500)])
    p = {"date": "2026-10-07", "title": "Electronics Components", "quotes": [q1, q2, quote("C", [("J", 1, 30000), ("K", 1, 10)])]}
    out = build_permission(p, profile, str(tmp_path / "h.docx"))
    t = text(out)
    for line in ("Jetson", "Case", "SSD"):
        assert f"\n{line}\n" in f"\n{t}\n"
    comp = table_rows(out, 1)
    assert [r[0] for r in comp[1:4]] == ["Item 1", "Item 2", "Item 3"]
    assert comp[2][2] == "" and comp[3][2] == ""               # BETA quoted fewer items


def test_duplicate_subject_prefix_is_removed(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "Permission to Purchase for Electronics Components",
         "quotes": [quote("A", [("x", 1, 100)])]}
    t = text(build_permission(p, profile, str(tmp_path / "i.docx")))
    assert "Permission to purchase for Electronics Components" in t
    assert t.lower().count("permission to purchase for") == 1


def test_special_characters_and_long_text(tmp_path, profile):
    desc = 'Motor <250W> & "Encoder" – 100RPM ™ ₹ ' + "very long description " * 20
    p = {"date": "2026-10-07", "title": "R&D <test> “quotes”", "quotes": [quote("A & B <Pvt> Ltd", [(desc, 1, 100)])]}
    out = build_permission(p, profile, str(tmp_path / "j.docx"))
    assert desc.strip() in text(out) or desc.strip() in table_rows(out, 0)[1][1]
    Document(out)                                                  # still a valid file


def test_thirty_items(tmp_path, profile):
    q = quote("A", [(f"Item {i}", i, 10 * i) for i in range(1, 31)])
    p = {"date": "2026-10-07", "title": "Many", "quotes": [q, quote("B", [("x", 1, 99999)]), quote("C", [("y", 1, 99999)])],
         "list_items_in_subject": False}
    out = build_permission(p, profile, str(tmp_path / "k.docx"))
    rows = table_rows(out, 0)
    assert rows[30][0] == "30" and money(rows[-1][1]) == quote_totals(q)["grand"]


def test_financial_year_in_reference(tmp_path, profile):
    for d, fy in (("2026-03-31", "2025-26"), ("2026-04-01", "2026-27")):
        p = {"date": d, "title": "x", "quotes": [quote("A", [("x", 1, 100)])]}
        assert f"/{fy}/" in text(build_permission(p, profile, str(tmp_path / "l.docx"))).splitlines()[0]


def test_incubation_letter(tmp_path):
    inc = store.all_("projects", "id ASC")[1]
    p = {"date": "2026-07-16", "title": "Claude Pro Subscription for One Month",
         "quotes": [{"vendor": {"name": "Anthropic"}, "items": [{"description": "Claude Pro", "qty": 1, "rate": 2033.05, "gst": 0, "period": "1 month"}]}]}
    out = build_permission(p, inc, str(tmp_path / "m.docx"))
    t = text(out)
    assert "PrithviX" in t and "Dr. Darshita Shah" in t and "Dr. Mehul Naik" in t and "NU/IT/Purchase" not in t
    assert table_rows(out, 0)[-1][1] == "₹ 2,033.05"


# ---------------------------------------------------------------- purchase order
def test_po_with_audit_page(tmp_path, profile):
    q = quote("DAZZLE ROBOTICS PVT. LTD.", [("Motor", 4, 3912)], quote_no="1103100", quote_date="2025-10-26")
    out = build_po({"date": "2025-12-20", "title": "Motor", "quotes": [q]}, profile, str(tmp_path / "po.docx"), BUDGET)
    t = text(out)
    assert "To\nDAZZLE ROBOTICS PVT. LTD.\nLine 1\nLine 2\nGST: 24AAAAA0000A1Z5" in t
    assert "Your Quotation Reference: 1103100 dated 26/10/2025" in t
    assert "2. Payment: After 15 days" in t and "THIS IS REQUIRED TO BE AUDITED" in t
    audit = table_rows(out, 1)[1]
    assert audit == ["1", "GA", "3,00,000", "2,59,572.10", "40,427.90", "18,464.64"]
    assert "Name of Authority: VP dated" in t


def test_po_without_audit_below_10k_and_without_quote_ref(tmp_path, profile):
    q = quote("A", [("x", 2, 3000)])
    out = build_po({"date": "2026-10-07", "title": "x", "quotes": [q]}, profile, str(tmp_path / "po2.docx"), BUDGET)
    t = text(out)
    assert "AUDITED" not in t and "Your Quotation Reference" not in t
    assert len(Document(out).tables) == 1


def test_rfq_letter(tmp_path, profile):
    p = {"date": "2026-10-07", "title": "Motor drivers", "quotes": [quote("A", [("Sabertooth", 2, 1)])]}
    out = build_rfq(p, profile, {"name": "Prayog", "address": "Ranchi"}, str(tmp_path / "rfq.docx"))
    t = text(out)
    assert "Request for quotation for Motor drivers" in t and "PRAYOG" in t
    rows = table_rows(out, 0)
    assert rows[1][1] == "Sabertooth" and rows[1][2] == "" and rows[1][3] == "2"   # no prices to the vendor


# ---------------------------------------------------------------- vouchers (Excel)
def cell(path, ref, sheet=None):
    wb = load_workbook(path)
    return (wb[sheet] if sheet else wb.worksheets[0])[ref].value


@pytest.mark.parametrize("n", [1, 3, 10, 11, 25])
def test_advance_voucher_rows(tmp_path, profile, n):
    items = [{"description": f"Item {i}", "qty": 2, "rate": 100 * i} for i in range(1, n + 1)]
    x = str(tmp_path / "adv.xlsx")
    build_advance_voucher({"date": "2026-10-07", "items": items}, profile, BUDGET, x, "x.pdf")
    total = sum(2 * 100 * i for i in range(1, n + 1))
    wb = load_workbook(x)
    ws = wb["Sheet1"]
    assert ws["G9"].value == total and ws["B10"].value.endswith("only")
    assert [ws[f"B{34 + i}"].value for i in range(n)] == [f"Item {i}" for i in range(1, n + 1)]
    assert ws[f"H{34 + n}"].value == total and ws[f"B{34 + n}"].value == "Grand Total"
    assert ws[f"B{35 + n}"].value is None                            # nothing left over below
    assert ws["E27"].value == round(BUDGET["provision"] - BUDGET["utilized"] - total, 2)
    assert ws["A1"].value == "☑" and ws["G13"].value.day == 6        # +30 days


@pytest.mark.parametrize("n", [1, 4, 30])
def test_advance_adjustment(tmp_path, profile, n):
    bills = [{"biller": f"Shop {i}", "bill_no": f"B/{i}", "bill_date": "2026-10-01", "gst": 10, "items": "x", "amount": 100.5 * i}
             for i in range(1, n + 1)]
    x = str(tmp_path / "adj.xlsx")
    build_advance_adjustment({"date": "2026-10-07", "bills": bills, "remaining_adjustment": 50}, profile, BUDGET, x, "x.pdf")
    ws = load_workbook(x)["Sheet2"]
    total = round(sum(b["amount"] for b in bills), 2)
    assert ws["F7"].value == total and ws["G32"].value == total and ws["E25"].value == round(total)
    assert ws["E27"].value == round(300000 - (259572.10 + 50 + round(total)), 2)
    assert ws[f"B{33 + n}"].value == f"Shop {n}" and ws[f"B{34 + n}"].value is None
    assert ws["A31"].value.startswith("NU/IT/KISAN KAWACH/2627/Advance Adjustment Voucher/")
    assert len(load_workbook(x).sheetnames) == 1


def test_cash_voucher(tmp_path, profile):
    bills = [{"biller": "HYDRO", "bill_no": "T/29", "bill_date": "2026-06-30", "items": "Wheel hub", "amount": 2950}]
    x = str(tmp_path / "cv.xlsx")
    build_cash_voucher({"date": "2026-10-07", "bills": bills, "permission_ref": "Permission dated 27-06-2026"}, profile, BUDGET, x, "x.pdf")
    ws = load_workbook(x)["Sheet1"]
    assert ws["D5"].value == "CASH VOUCHER" and ws["G9"].value == 2950
    assert ws["B34"].value == "HYDRO – Wheel hub" and ws["H35"].value == 2950
    assert ws["G13"].value == "Permission dated 27-06-2026"


# ---------------------------------------------------------------- post-facto and waiver
from gen_word import build_waiver  # noqa: E402


def test_post_facto_permission(tmp_path, profile):
    q = quote("GROWIT INDIA PVT LTD", [("GROWIT SOIL GURU PRO", 1, 30508)], round_off=0.56)
    p = {"date": "2026-01-01", "title": "GROWIT SOIL GURU PRO (WITHOUT SUBCRIPTION)", "quotes": [q], "post_facto": True,
         "justification": "Sensors were running out of stock."}
    out = build_permission(p, profile, str(tmp_path / "pf.docx"))
    t = text(out)
    assert "Justification for purchase without prior Permission:" in t and "Sensors were running out of stock." in t
    assert "post facto approval" in t and "Total cost of purchase: ₹ 36,000.00" in t
    assert "Supplier Details: GROWIT INDIA PVT LTD" in t and "Through" in t          # 36k -> Director route
    assert table_rows(out, 0)[-1] == ["Grand Total", "Rs. 36,000.00"]
    assert " Statement" in t


def test_post_facto_default_justification(tmp_path, profile):
    p = {"date": "2026-01-01", "title": "x", "quotes": [quote("A", [("x", 1, 100)])], "post_facto": True}
    assert "personal funds" in text(build_permission(p, profile, str(tmp_path / "pf2.docx")))


def test_waiver(tmp_path, profile):
    rec = {"date": "2026-01-15", "waiver_reason": "Urgent need during integration.",
           "bills": [{"biller": "Growit India Private Limited", "items": "Soil Guru Pro", "amount": 36000, "bill_no": "CKD-00049", "bill_date": "2025-12-27"},
                     {"biller": "Robu", "items": "Motor driver", "amount": 28695.99, "bill_no": "INV/1", "bill_date": "2025-12-08"}]}
    out = build_waiver(rec, profile, str(tmp_path / "w.docx"))
    t = text(out)
    assert "Date :   15  / 1  / 26" in t and "Urgent need during integration." in t and "Vice President" in t
    rows = table_rows(out, 0)
    assert len(rows) == 3
    assert rows[1][1] == "GROWIT INDIA PRIVATE LIMITED" and rows[1][3] == "36,000" and "27/12/25" in rows[1][4]
    assert rows[2][3] == "28,695.99"
